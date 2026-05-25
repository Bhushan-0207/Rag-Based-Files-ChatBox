import os
import pandas as pd

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def parse_pdf(path):

    reader = PdfReader(path)

    documents = []

    for page_num, page in enumerate(reader.pages):

        text = page.extract_text()

        if text:

            documents.append({
                "content": text,
                "metadata": {
                    "page": page_num + 1
                }
            })

    return documents


def parse_docx(path):

    doc = Document(path)

    text = "\n".join(
        para.text
        for para in doc.paragraphs
    )

    return [{
        "content": text,
        "metadata": {
            "section": "document"
        }
    }]


def parse_xlsx(path):

    excel_file = pd.ExcelFile(path)

    documents = []

    for sheet in excel_file.sheet_names:

        df = excel_file.parse(sheet)

        text = df.to_string()

        documents.append({
            "content": text,
            "metadata": {
                "sheet": sheet
            }
        })

    return documents


def parse_pptx(path):

    prs = Presentation(path)

    documents = []

    for slide_num, slide in enumerate(prs.slides):

        slide_text = ""

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                slide_text += shape.text + "\n"

        documents.append({
            "content": slide_text,
            "metadata": {
                "slide": slide_num + 1
            }
        })

    return documents


def extract_text(file_path):

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return parse_pdf(file_path)

    elif ext == ".docx":
        return parse_docx(file_path)

    elif ext == ".xlsx":
        return parse_xlsx(file_path)

    elif ext == ".pptx":
        return parse_pptx(file_path)

    else:
        raise ValueError("Unsupported file type")